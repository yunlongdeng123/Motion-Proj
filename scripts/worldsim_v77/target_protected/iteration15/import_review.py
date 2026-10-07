"""只读导入新版人工工作簿；保留空白、原备注与单元格出处。"""
from pathlib import Path
import argparse, json, shutil
from openpyxl import load_workbook


def main(source, dest):
    dest.mkdir(parents=True, exist_ok=True)
    archived = dest/source.name
    if source.resolve()!=archived.resolve():
        shutil.copy2(source, archived)
    book = load_workbook(archived, read_only=True, data_only=False)
    raw = []; records = []
    for sheet in book:
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        raw.append({'sheet': sheet.title, 'rows': rows})
        version = 'r47' if sheet.title.startswith('r47') else 'r46' if sheet.title.startswith('r46') else None
        if version:
            for number, row in enumerate(rows, 1):
                label = str(row[0] or '')
                if not label.startswith('A'):
                    continue
                cid = label.split()[0].split('·')[0].strip()
                score = row[1]
                if score not in (None, 0, 1, 2):
                    raise ValueError(f'未知人工分: {sheet.title}!B{number}={score}')
                records.append({'version': version, 'case_id': cid, 'label': label,
                    'human_score': score, 'RGB_prior_sufficient': row[2], 'OCC_prior_sufficient': row[3],
                    'note': row[5], 'source': str(source), 'sheet': sheet.title,
                    'range': f'A{number}:F{number}'})
    book.close()
    assert len([r for r in records if r['version']=='r47'])==8
    assert len([r for r in records if r['version']=='r46'])==8
    result = {'source': str(source), 'archived_original': str(archived),
        'source_modified': False, 'sheets': raw, 'records': records,
        'r47_unfilled': [r['case_id'] for r in records if r['version']=='r47' and r['human_score'] is None],
        'scope': '用户人工记录；空白不补分，旧r46与新r47分开；不代填时序通过率'}
    (dest/'human_review.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8', newline='\n')
    print(json.dumps({'sheets': len(raw), 'records': len(records), 'r47_unfilled': result['r47_unfilled']}, ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source', type=Path, required=True);p.add_argument('--dest', type=Path, required=True)
    a=p.parse_args();main(a.source, a.dest)
